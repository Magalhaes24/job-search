"""Lê as vagas do Excel (folha 'Dados completos') e permite filtrá-las."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
XLSX_PADRAO = RAIZ.parent / "Estagios_Janeiro2027.xlsx"

COLUNAS = {
    "id": "ID", "empresa": "Empresa", "role": "Role", "area": "Tipo / área",
    "tipo_empresa": "Tipo de empresa", "pais": "País", "cidade": "Cidade",
    "modalidade": "Modalidade", "inicio": "Início", "duracao": "Duração",
    "descricao": "Descrição (o que vais fazer)", "requisitos": "Requirements",
    "idiomas": "Idiomas", "elegibilidade": "Elegibilidade (estudante / recém-licenciada)",
    "deadline": "Deadline", "dificuldade": "Dificuldade de ser aceite",
    "fit": "Fit com PE/Finance (1-5)", "lingua": "Língua-chave",
    "match": "Match com o CV (0-100)", "prob": "Prob. estimada de aceitação",
    "chance": "Chance", "estado": "Estado", "veredicto": "Veredicto",
    "notas": "Notas / riscos", "link": "Link", "tipo_link": "Tipo de link",
    "qualidade": "Qualidade da informação", "ativa": "Ativa", "rank": "Rank",
}


@dataclass
class Vaga:
    id: int
    empresa: str
    role: str
    area: str
    tipo_empresa: str
    pais: str
    cidade: str
    modalidade: str
    inicio: str
    duracao: str
    descricao: str
    requisitos: str
    idiomas: str
    elegibilidade: str
    deadline: str
    dificuldade: str
    fit: float
    lingua: str
    match: float
    prob: float
    chance: str
    estado: str
    veredicto: str
    notas: str
    link: str
    tipo_link: str
    qualidade: str
    ativa: int
    rank: int

    @property
    def espontanea(self) -> bool:
        t = f"{self.area} {self.role}".lower()
        return "espontânea" in t or "via a explorar" in self.estado.lower()

    @property
    def cidade_valida(self) -> bool:
        return bool(self.cidade) and "confirmar" not in self.cidade.lower()


def _txt(v) -> str:
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%d-%m-%Y")
    return str(v).strip()


def _num(v, padrao=0.0) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return padrao


def carregar(caminho=None) -> list[Vaga]:
    p = Path(caminho) if caminho else XLSX_PADRAO
    if not p.exists():
        raise FileNotFoundError(
            f"Não encontrei o Excel em {p}. Usa --xlsx para indicar o caminho."
        )
    ws = openpyxl.load_workbook(p, data_only=True)["Dados completos"]
    cab = {c.value: i for i, c in enumerate(ws[3]) if c.value}
    vagas = []
    for linha in ws.iter_rows(min_row=4, values_only=True):
        if not linha or linha[cab["ID"]] is None:
            continue
        g = lambda k: linha[cab[COLUNAS[k]]]
        vagas.append(Vaga(
            id=int(g("id")), empresa=_txt(g("empresa")), role=_txt(g("role")),
            area=_txt(g("area")), tipo_empresa=_txt(g("tipo_empresa")),
            pais=_txt(g("pais")), cidade=_txt(g("cidade")),
            modalidade=_txt(g("modalidade")), inicio=_txt(g("inicio")),
            duracao=_txt(g("duracao")), descricao=_txt(g("descricao")),
            requisitos=_txt(g("requisitos")), idiomas=_txt(g("idiomas")),
            elegibilidade=_txt(g("elegibilidade")), deadline=_txt(g("deadline")),
            dificuldade=_txt(g("dificuldade")), fit=_num(g("fit")),
            lingua=_txt(g("lingua")) or "EN", match=_num(g("match")),
            prob=_num(g("prob")), chance=_txt(g("chance")), estado=_txt(g("estado")),
            veredicto=_txt(g("veredicto")), notas=_txt(g("notas")),
            link=_txt(g("link")), tipo_link=_txt(g("tipo_link")),
            qualidade=_txt(g("qualidade")), ativa=int(_num(g("ativa"), 1)),
            rank=int(_num(g("rank"), 9999)),
        ))
    return vagas


def filtrar(vagas, ids=None, empresa=None, top=None, prioridades=False,
            pais=None, area=None, incluir_inativas=False):
    sel = list(vagas)
    if not incluir_inativas and not ids:
        sel = [v for v in sel if v.ativa == 1]
    if ids:
        conj = set(ids)
        sel = [v for v in vagas if v.id in conj]
    if empresa:
        sel = [v for v in sel if empresa.lower() in v.empresa.lower()]
    if pais:
        sel = [v for v in sel if pais.lower() in v.pais.lower()]
    if area:
        sel = [v for v in sel if area.lower() in f"{v.area} {v.tipo_empresa}".lower()]
    if prioridades:
        sel = [v for v in sel if v.veredicto == "Prioridade"]
    sel.sort(key=lambda v: (v.rank, v.id))
    if top:
        sel = sel[:top]
    return sel
