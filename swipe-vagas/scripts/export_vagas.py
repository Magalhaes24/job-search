"""Exporta as vagas do Excel (folha 'Dados completos') para src/data/vagas.json.

Uso:  python scripts/export_vagas.py [caminho/para/Estagios_Janeiro2027.xlsx]
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ.parent / "Estagios_Janeiro2027.xlsx"
SAIDA = RAIZ / "src" / "data" / "vagas.json"

CAMPOS = {
    "id": "ID", "empresa": "Empresa", "role": "Role", "area": "Tipo / área",
    "tipoEmpresa": "Tipo de empresa", "pais": "País", "cidade": "Cidade",
    "modalidade": "Modalidade", "inicio": "Início", "duracao": "Duração",
    "descricao": "Descrição (o que vais fazer)", "requisitos": "Requirements",
    "idiomas": "Idiomas", "elegibilidade": "Elegibilidade (estudante / recém-licenciada)",
    "deadline": "Deadline", "salarioIndicado": "Salário indicado na oferta (como indicado)",
    "salMin": "Salário bruto est. mín (€/mês)", "salMax": "Salário bruto est. máx (€/mês)",
    "salMed": "Salário bruto est. médio (€/mês)", "fonteSalario": "Fonte / confiança do salário bruto",
    "dificuldade": "Dificuldade de ser aceite", "fit": "Fit com PE/Finance (1-5)",
    "facilidade": "Facilidade (1-5)", "score": "Score (60% fit + 40% facilidade)",
    "lingua": "Língua-chave", "exigLingua": "Exigência da língua (0-1)",
    "fIdioma": "Idioma (0-1)", "fExperiencia": "Experiência (0-1)",
    "fElegibilidade": "Elegibilidade (0-1)", "fPerfil": "Perfil / skills (0-1)",
    "match": "Match com o CV (0-100)", "probBase": "Prob. base (dificuldade)",
    "prob": "Prob. estimada de aceitação", "chance": "Chance", "porque": "Porquê (CV)",
    "estado": "Estado", "veredicto": "Veredicto", "notas": "Notas / riscos",
    "link": "Link", "tipoLink": "Tipo de link", "qualidade": "Qualidade da informação",
    "ativa": "Ativa", "pEfetiva": "P efetiva", "indice": "Índice (P × interesse)", "rank": "Rank",
}
NUMEROS = {"id", "salMin", "salMax", "salMed", "dificuldade_n", "fit", "facilidade", "score",
           "exigLingua", "fIdioma", "fExperiencia", "fElegibilidade", "fPerfil", "match",
           "probBase", "prob", "ativa", "pEfetiva", "indice", "rank"}


def limpar(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%d-%m-%Y")
    return str(v).strip()


def resumo(desc: str, limite=190) -> str:
    desc = re.sub(r"\s+", " ", desc).strip()
    if len(desc) <= limite:
        return desc
    corte = desc[:limite]
    ponto = max(corte.rfind(". "), corte.rfind("; "))
    if ponto > 80:
        return corte[: ponto + 1]
    return corte[: corte.rfind(" ")].rstrip(",;:") + "…"


ws = openpyxl.load_workbook(XLSX, data_only=True)["Dados completos"]
cab = {c.value: i for i, c in enumerate(ws[3]) if c.value}
vagas = []
for linha in ws.iter_rows(min_row=4, values_only=True):
    if not linha or linha[cab["ID"]] is None:
        continue
    v = {}
    for k, col in CAMPOS.items():
        bruto = linha[cab[col]]
        if k in NUMEROS:
            try:
                v[k] = round(float(bruto), 5)
            except (TypeError, ValueError):
                v[k] = None
        else:
            v[k] = limpar(bruto)
    v["id"] = int(v["id"])
    v["ativa"] = int(v["ativa"] if v["ativa"] is not None else 1)
    v["rank"] = int(v["rank"]) if v["rank"] is not None else 9999
    v["resumo"] = resumo(v["descricao"])
    vagas.append(v)

SAIDA.parent.mkdir(parents=True, exist_ok=True)
SAIDA.write_text(json.dumps(vagas, ensure_ascii=False, indent=None, separators=(",", ":")), encoding="utf-8")
print(f"{len(vagas)} vagas -> {SAIDA} ({SAIDA.stat().st_size // 1024} KB)")
